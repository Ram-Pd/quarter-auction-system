"""
Excel import of quarters + their auctions (auctions/importers.py).

The sample workbooks are built in memory by sample_data/make_samples.py with a
bidding date a week from today, so these tests never go stale. One test also
checks the committed .xlsx files in sample_data/ have the right shape.
"""
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from accounts.models import User
from applications.tests.factories import make_staff
from auctions.importers import EXPECTED_HEADERS, import_quarters
from auctions.models import Auction, Quarter
from sample_data import make_samples

DAY = date.today() + timedelta(days=7)


def good_row(quarter_no="B-214", **changes):
    row = dict(zip(EXPECTED_HEADERS, ["B-214", "Sector 4", "Near market", "Type B", "2nd", 750, 45000, 9000,
                                      DAY, time(10, 0), 5]))
    row["Quarter No"] = quarter_no
    row.update(changes)
    return [row[h] for h in EXPECTED_HEADERS]


def xlsx(rows, headers=EXPECTED_HEADERS):
    """Build an in-memory .xlsx file from a list of rows."""
    wb = Workbook()
    ws = wb.active
    if headers is not None:
        ws.append(headers)
    for row in rows:
        ws.append(row)
    return saved(wb)


def saved(workbook):
    buf = BytesIO()
    workbook.save(buf)
    buf.seek(0)
    return buf


def messages(result):
    return [str(e) for e in result.errors]


def ist(day, hour, minute):
    return timezone.make_aware(datetime.combine(day, time(hour, minute)), timezone.get_default_timezone())


class SampleFileTests(TestCase):
    """The four sample workbooks behave as documented in sample_data/README.md."""

    def test_valid_sample_creates_eight_quarters_with_auctions(self):
        result = import_quarters(saved(make_samples.build_valid(DAY)))
        self.assertEqual(messages(result), [])
        self.assertEqual(result.created, 8)
        self.assertEqual(Quarter.objects.count(), 8)
        self.assertEqual(Auction.objects.count(), 8)
        q = Quarter.objects.get(quarter_no="A-021")
        self.assertEqual((q.quarter_type, q.floor, q.area_sqft, q.base_rate, q.emd_amount),
                         ("A", "Ground", Decimal("980"), Decimal("62000"), Decimal("12000")))
        self.assertEqual(q.status, Quarter.Status.IN_AUCTION)
        auction = q.auctions.get()
        self.assertEqual(auction.status, Auction.Status.SCHEDULED)
        self.assertEqual(auction.start_time, ist(DAY, 11, 0))          # read as India time
        self.assertEqual(auction.duration_minutes, 10)
        self.assertEqual(auction.end_time, ist(DAY, 11, 10))
        # Blank duration -> default.
        self.assertEqual(Auction.objects.get(quarter__quarter_no="D-407").duration_minutes,
                         settings.DEFAULT_AUCTION_DURATION_MINUTES)

    def test_invalid_sample_reports_every_error_and_saves_nothing(self):
        result = import_quarters(saved(make_samples.build_invalid(DAY)))
        found = {(e.row, e.column) for e in result.errors}
        self.assertEqual(found, {
            (3, "Quarter No"),          # duplicate B-214
            (4, "Area (sq ft)"),        # "abc"
            (5, "Base Rate (Rs.)"),     # negative
            (6, "Quarter No"),          # missing
            (7, "EMD Amount (Rs.)"),    # 0
            (8, "Bidding Date"),        # 31-02-2026
            (9, "Start Time"),          # yesterday
            (10, "Start Time"),         # 25:00
            (10, "Duration (min)"),     # 0
        })
        self.assertEqual(result.created, 0)
        self.assertEqual(Quarter.objects.count(), 0)  # row 2 was valid but NOT saved
        self.assertEqual(Auction.objects.count(), 0)

    def test_bad_header_sample_is_rejected(self):
        result = import_quarters(saved(make_samples.build_bad_header(DAY)))
        self.assertEqual(messages(result), ["Row 1, Column A: expected header 'Quarter No', found 'Quarter Number'"])

    def test_edge_case_sample_imports_cleanly(self):
        result = import_quarters(saved(make_samples.build_edge_cases(DAY)))
        self.assertEqual(messages(result), [])
        self.assertEqual(result.created, 4)
        e101 = Quarter.objects.get(quarter_no="E-101")
        self.assertEqual((e101.sector, e101.location_block, e101.quarter_type), ("Sector 5", "Behind hospital", "A"))
        self.assertEqual((e101.area_sqft, e101.base_rate, e101.emd_amount),
                         (Decimal("1050"), Decimal("75000"), Decimal("15000")))
        self.assertEqual(e101.auctions.get().start_time, ist(DAY, 9, 30))       # text date + text time
        e102 = Quarter.objects.get(quarter_no="E-102")
        self.assertEqual(e102.base_rate, Decimal("72500.50"))
        self.assertEqual(e102.auctions.get().start_time, ist(DAY, 14, 15))      # "2:15 PM"
        self.assertEqual(e102.auctions.get().duration_minutes, settings.DEFAULT_AUCTION_DURATION_MINUTES)
        self.assertEqual(Quarter.objects.get(quarter_no="E-104").auctions.get().start_time, ist(DAY, 16, 45))

    def test_importing_same_file_twice_rejects_all_as_existing(self):
        import_quarters(saved(make_samples.build_valid(DAY)))
        result = import_quarters(saved(make_samples.build_valid(DAY)))
        self.assertEqual(len(result.errors), 8)
        self.assertTrue(all("already exists in the database" in m for m in messages(result)))

    def test_committed_sample_files_have_the_expected_headers(self):
        for name in ["quarters_sample", "quarters_invalid", "quarters_edge_cases"]:
            with self.subTest(file=name):
                sheet = load_workbook(settings.BASE_DIR / "sample_data" / f"{name}.xlsx").worksheets[0]
                self.assertEqual([c.value for c in sheet[1]], EXPECTED_HEADERS)


class ValidationRuleTests(TestCase):
    def test_duplicate_against_database_case_insensitive(self):
        import_quarters(xlsx([good_row()]))
        result = import_quarters(xlsx([good_row("b-214", **{"Start Time": time(11, 0)})]))
        self.assertEqual(messages(result), ["Row 2, Quarter No: 'B-214' already exists in the database."])

    def test_numeric_quarter_no_stored_as_text(self):
        result = import_quarters(xlsx([good_row(214)]))
        self.assertTrue(result.ok)
        self.assertTrue(Quarter.objects.filter(quarter_no="214").exists())

    def test_every_error_in_a_row_is_reported(self):
        row = ["", "", "", "Type Z", "", "x", 0, -1, "someday", "noon", "five"]
        result = import_quarters(xlsx([row]))
        self.assertEqual([e.column for e in result.errors], [
            "Quarter No", "Sector", "Location / Block", "Floor", "Type", "Area (sq ft)", "Base Rate (Rs.)",
            "EMD Amount (Rs.)", "Bidding Date", "Start Time", "Duration (min)"])

    def test_date_and_time_formats(self):
        accepted = [
            (DAY.strftime("%d-%m-%Y"), "10:30"), (DAY.strftime("%d/%m/%Y"), "10:30 am"),
            (DAY.strftime("%Y-%m-%d"), "10.30"), (DAY.strftime("%d %b %Y"), 0.4375),   # 0.4375 day = 10:30
            (datetime.combine(DAY, time()), datetime(1899, 12, 30, 10, 30)),
        ]
        for i, (bidding_date, start) in enumerate(accepted):
            with self.subTest(date=bidding_date, time=start):
                result = import_quarters(xlsx([good_row(f"T-{i}", **{"Bidding Date": bidding_date,
                                                                       "Start Time": start})]))
                self.assertEqual(messages(result), [])
                self.assertEqual(Auction.objects.get(quarter__quarter_no=f"T-{i}").start_time, ist(DAY, 10, 30))

    def test_start_must_be_in_the_future(self):
        now = timezone.now()
        today = timezone.localdate()
        one_minute_ago = timezone.localtime(now - timedelta(minutes=1))
        result = import_quarters(xlsx([good_row(**{"Bidding Date": today, "Start Time": one_minute_ago.time()})]),
                                 now=now)
        self.assertEqual([e.column for e in result.errors], ["Start Time"])
        self.assertIn("must be in the future", result.errors[0].message)

    def test_duration_rules(self):
        cases = [("", True), (None, True), (7, True), ("10", True), (7.0, True), (0, False),
                 (-5, False), (2.5, False), ("abc", False), (1441, False)]
        for i, (value, ok) in enumerate(cases):
            with self.subTest(duration=value):
                result = import_quarters(xlsx([good_row(f"D-{i}", **{"Duration (min)": value})]))
                self.assertEqual(result.ok, ok, messages(result))

    def test_whitespace_commas_and_type_spelling(self):
        row = good_row("  c-305 ", **{"Sector": "  Sector   2 ", "Type": "type c", "Base Rate (Rs.)": "45,500",
                                      "EMD Amount (Rs.)": " 9,000.50 "})
        result = import_quarters(xlsx([row, [None] * 11, ["", " ", None] + [""] * 8]))
        self.assertEqual(messages(result), [])
        q = Quarter.objects.get(quarter_no="C-305")
        self.assertEqual((q.sector, q.quarter_type, q.base_rate, q.emd_amount),
                         ("Sector 2", "C", Decimal("45500"), Decimal("9000.50")))

    def test_too_many_decimal_places_and_too_large(self):
        result = import_quarters(xlsx([good_row("X-1", **{"Area (sq ft)": "750.123"}),
                                       good_row("X-2", **{"Base Rate (Rs.)": "99999999999"})]))
        self.assertEqual(messages(result), ["Row 2, Area (sq ft): can have at most 2 decimal places",
                                            "Row 3, Base Rate (Rs.): is too large"])

    def test_missing_extra_and_misordered_headers(self):
        headers = ["Sector", "Quarter No"] + EXPECTED_HEADERS[2:10] + [None, "Notes"]
        result = import_quarters(xlsx([], headers=headers))
        self.assertEqual(messages(result), [
            "Row 1, Column A: expected header 'Quarter No', found 'Sector'",
            "Row 1, Column B: expected header 'Sector', found 'Quarter No'",
            "Row 1, Column K: missing header 'Duration (min)'",
            "Row 1, Column L: unexpected extra column 'Notes'",
        ])

    def test_header_only_file_empty_sheet_and_not_excel(self):
        self.assertEqual(messages(import_quarters(xlsx([]))), ["The file has a header but no data rows."])
        self.assertEqual(messages(import_quarters(xlsx([], headers=None))), ["The sheet is empty."])
        self.assertEqual(messages(import_quarters(BytesIO(b"not a spreadsheet"))),
                         ["File is not a valid .xlsx Excel workbook."])


class ImportAdminPageTests(TestCase):
    def setUp(self):
        self.url = reverse("admin:auctions_quarter_import")
        self.auction_admin = make_staff("aa@example.com", "Auction Admin")

    def upload(self, workbook, name="quarters.xlsx", **kwargs):
        data = saved(workbook).read() if not isinstance(workbook, bytes) else workbook
        return self.client.post(self.url, {"file": SimpleUploadedFile(name, data)}, **kwargs)

    def test_requires_login(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 302)
        self.assertIn("/admin/login/", response["Location"])

    def test_verification_officer_cannot_import(self):
        self.client.force_login(make_staff("vo@example.com", "Verification Officer"))
        self.assertEqual(self.client.get(self.url).status_code, 403)

    def test_changelist_has_import_button(self):
        self.client.force_login(self.auction_admin)
        response = self.client.get(reverse("admin:auctions_quarter_changelist"))
        self.assertContains(response, "Import from Excel")

    def test_auction_admin_can_upload(self):
        self.client.force_login(self.auction_admin)
        response = self.upload(make_samples.build_valid(DAY), follow=True)
        self.assertRedirects(response, reverse("admin:auctions_auction_changelist"))
        self.assertContains(response, "Imported 8 quarter(s), each with its scheduled auction.")
        self.assertEqual(Auction.objects.count(), 8)

    def test_invalid_upload_shows_error_table(self):
        self.client.force_login(self.auction_admin)
        response = self.upload(make_samples.build_invalid(DAY))
        self.assertContains(response, "Nothing was imported. The file has 9 errors.")
        self.assertContains(response, "must be a number (got &#x27;abc&#x27;)")
        self.assertEqual(Quarter.objects.count(), 0)

    def test_wrong_extension_rejected(self):
        self.client.force_login(self.auction_admin)
        response = self.upload(b"a,b,c", name="quarters.csv")
        self.assertContains(response, "File extension")


@override_settings(DEFAULT_AUCTION_DURATION_MINUTES=12)
class DefaultDurationSettingTests(TestCase):
    def test_blank_duration_uses_setting(self):
        import_quarters(xlsx([good_row(**{"Duration (min)": None})]))
        self.assertEqual(Auction.objects.get().duration_minutes, 12)


class SuperuserImportTests(TestCase):
    def test_superuser_can_open_import_page(self):
        self.client.force_login(User.objects.create_superuser("boss@example.com", "x", full_name="Boss"))
        self.assertEqual(self.client.get(reverse("admin:auctions_quarter_import")).status_code, 200)
