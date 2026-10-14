"""
Excel import of quarters AND their auctions (one bidding slot per quarter).

Two phases:
  1. Validate the header and EVERY row, collecting all errors (nothing saved).
  2. Only if there are zero errors, save all rows inside one transaction:
     each row creates a Quarter and its Auction (status SCHEDULED).

Dates and times are read as India time (settings.TIME_ZONE = Asia/Kolkata).
Both real Excel date/time cells and typed text are accepted.
"""
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from django.conf import settings
from django.db import IntegrityError, transaction
from django.db.models.functions import Upper
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .models import Auction, Quarter

EXPECTED_HEADERS = [
    "Quarter No", "Sector", "Location / Block", "Type", "Floor", "Area (sq ft)",
    "Base Rate (Rs.)", "EMD Amount (Rs.)", "Bidding Date", "Start Time", "Duration (min)",
]
MAX_DATA_ROWS = 2000            # protects the server from huge uploads
MAX_DURATION_MINUTES = 24 * 60  # one auction slot can't be longer than a day

# Accepts "Type B", "type b", "B" -> "B"
TYPE_LOOKUP = {}
for _value, _label in Quarter.QuarterType.choices:
    TYPE_LOOKUP[_value.lower()] = _value
    TYPE_LOOKUP[_label.lower()] = _value

# Text dates we understand, tried in this order. Day comes before month (Indian style).
DATE_FORMATS = ["%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d-%b-%Y", "%d %b %Y", "%d %B %Y"]
TIME_FORMATS = ["%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M:%S %p", "%I %p", "%H.%M"]


@dataclass
class RowError:
    row: int | None      # Excel row number (1 = header); None = whole-file problem
    column: str | None   # header name, e.g. "Area (sq ft)"
    message: str

    def __str__(self):
        where = []
        if self.row is not None:
            where.append(f"Row {self.row}")
        if self.column:
            where.append(self.column)
        return f"{', '.join(where)}: {self.message}" if where else self.message


@dataclass
class ImportResult:
    created: int = 0
    errors: list[RowError] = field(default_factory=list)

    @property
    def ok(self):
        return not self.errors


@dataclass
class ValidRow:
    """One fully checked row, ready to be saved."""
    quarter: Quarter
    start_time: datetime
    duration_minutes: int


# ---------- cell helpers: each returns (value, None) or (None, "error message") ----------

def _clean_text(value):
    """Cell -> stripped string ('' for blank). 214.0 -> '214'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())  # trims and collapses inner whitespace


def _parse_positive_decimal(value, max_digits, decimal_places=2):
    if isinstance(value, bool):  # True/False are ints in Python; reject them
        return None, "must be a number"
    if isinstance(value, (int, float)):
        text = str(value)
    else:
        text = _clean_text(value).replace(",", "")  # "4,500" -> "4500"
    if text == "":
        return None, "is required"
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None, f"must be a number (got '{_clean_text(value)}')"
    if not number.is_finite():
        return None, "must be a number"
    if number <= 0:
        return None, f"must be greater than 0 (got {_clean_text(value)})"
    if -number.as_tuple().exponent > decimal_places:
        return None, f"can have at most {decimal_places} decimal places"
    if number >= Decimal(10) ** (max_digits - decimal_places):
        return None, "is too large"
    return number, None


def _parse_date(value):
    if isinstance(value, datetime):          # a real Excel date cell
        return value.date(), None
    if isinstance(value, date):
        return value, None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # A date cell formatted as a plain number: Excel stores days since 1899-12-30.
        if 20000 < value < 100000:
            return from_excel(value).date(), None
        return None, f"is not a valid date (got {_clean_text(value)})"
    text = _clean_text(value)
    if not text:
        return None, "is required"
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date(), None
        except ValueError:
            continue
    return None, f"is not a valid date (got '{text}'). Use DD-MM-YYYY, e.g. 15-11-2026"


def _parse_time(value):
    if isinstance(value, datetime):          # some files store time cells as full date-times
        return value.time().replace(microsecond=0), None
    if isinstance(value, time):              # a real Excel time cell
        return value.replace(microsecond=0), None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # A time cell formatted as a number: a fraction of a day (0.4375 = 10:30).
        if 0 <= value < 1:
            seconds = round(value * 24 * 60 * 60)
            return (datetime.min + timedelta(seconds=seconds)).time(), None
        return None, f"is not a valid time (got {_clean_text(value)})"
    text = _clean_text(value).upper()
    if not text:
        return None, "is required"
    for fmt in TIME_FORMATS:
        try:
            return datetime.strptime(text, fmt).time(), None
        except ValueError:
            continue
    return None, f"is not a valid time (got '{_clean_text(value)}'). Use 24-hour HH:MM, e.g. 10:30"


def _parse_duration(value):
    """Blank -> the default; otherwise a whole number of minutes, 1 to 1440."""
    if isinstance(value, bool):
        return None, "must be a whole number of minutes"
    text = _clean_text(value)
    if text == "":
        return settings.DEFAULT_AUCTION_DURATION_MINUTES, None
    if isinstance(value, float) and not value.is_integer():
        return None, f"must be a whole number of minutes (got {text})"
    if not re.fullmatch(r"-?\d+", text):
        return None, f"must be a whole number of minutes (got '{text}')"
    minutes = int(text)
    if minutes <= 0:
        return None, f"must be greater than 0 (got {minutes})"
    if minutes > MAX_DURATION_MINUTES:
        return None, f"can be at most {MAX_DURATION_MINUTES} minutes (one day)"
    return minutes, None


# ---------- main entry point ----------

def import_quarters(uploaded_file, now=None) -> ImportResult:
    """Validate an .xlsx file and, if fully valid, create all quarters and their auctions."""
    result = ImportResult()
    now = now or timezone.now()

    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError):
        result.errors.append(RowError(None, None, "File is not a valid .xlsx Excel workbook."))
        return result

    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)

        # --- header ---
        header = next(rows, None)
        if header is None:
            result.errors.append(RowError(None, None, "The sheet is empty."))
            return result
        result.errors.extend(_check_header(header))
        if result.errors:
            return result  # columns can't be trusted, so don't read data rows

        # --- data rows ---
        valid_rows = []
        first_seen = {}  # quarter_no -> row where it first appeared
        data_rows = 0
        for row_no, cells in enumerate(rows, start=2):
            cells = list(cells) + [None] * (len(EXPECTED_HEADERS) - len(cells))
            if all(_clean_text(c) == "" for c in cells):
                continue  # skip blank rows
            data_rows += 1
            if data_rows > MAX_DATA_ROWS:
                result.errors.append(RowError(
                    None, None, f"Too many rows: the limit is {MAX_DATA_ROWS} quarters per file."))
                return result
            quarter_no, valid_row, errors = _validate_row(row_no, cells[:len(EXPECTED_HEADERS)], now)
            result.errors.extend(errors)
            # Duplicate checks run even if other columns in this row had errors,
            # so the admin sees every problem in one go.
            if quarter_no:
                if quarter_no in first_seen:
                    result.errors.append(RowError(
                        row_no, "Quarter No",
                        f"'{quarter_no}' is duplicated in this file "
                        f"(first seen at row {first_seen[quarter_no]})."))
                    continue
                first_seen[quarter_no] = row_no
            if valid_row is not None:
                valid_rows.append(valid_row)
    finally:
        workbook.close()

    if data_rows == 0 and not result.errors:
        result.errors.append(RowError(None, None, "The file has a header but no data rows."))

    # --- duplicates against the database ---
    existing = set(
        Quarter.objects.annotate(qn=Upper("quarter_no"))
        .filter(qn__in=list(first_seen))
        .values_list("qn", flat=True)
    )
    for quarter_no in sorted(existing, key=first_seen.get):
        result.errors.append(RowError(
            first_seen[quarter_no], "Quarter No", f"'{quarter_no}' already exists in the database."))

    if result.errors:
        result.errors.sort(key=lambda e: (e.row or 0))
        return result

    # --- save: all or nothing ---
    try:
        with transaction.atomic():
            for row in valid_rows:
                row.quarter.save()
                # Auction.save() sets end_time = start + duration and marks the quarter "In Auction".
                Auction.objects.create(quarter=row.quarter, start_time=row.start_time,
                                       duration_minutes=row.duration_minutes)
    except IntegrityError:
        # e.g. another admin imported the same quarter a moment ago
        result.errors.append(RowError(
            None, None, "Import cancelled: a quarter number was saved by someone else "
                        "during the import. Nothing was imported; please try again."))
        return result

    result.created = len(valid_rows)
    return result


def _check_header(header):
    errors = []
    header = [_clean_text(h) for h in header]
    # Ignore empty trailing columns (Excel often keeps formatting there).
    while header and header[-1] == "":
        header.pop()
    for i, expected in enumerate(EXPECTED_HEADERS):
        found = header[i] if i < len(header) else ""
        if found != expected:
            col = chr(ord("A") + i)
            errors.append(RowError(
                1, f"Column {col}",
                f"expected header '{expected}', found '{found}'" if found
                else f"missing header '{expected}'"))
    for i in range(len(EXPECTED_HEADERS), len(header)):
        if header[i]:
            errors.append(RowError(1, f"Column {chr(ord('A') + i)}",
                                   f"unexpected extra column '{header[i]}'"))
    return errors


def _validate_row(row_no, cells, now):
    """Return (quarter_no, ValidRow or None, list of RowError)."""
    errors = []
    values = dict(zip(EXPECTED_HEADERS, cells))

    def check(column, parsed):
        value, error = parsed
        if error:
            errors.append(RowError(row_no, column, error))
        return value

    def text(column, max_length):
        value = _clean_text(values[column])
        if value == "":
            errors.append(RowError(row_no, column, "is required"))
        elif len(value) > max_length:
            errors.append(RowError(row_no, column, f"must be at most {max_length} characters"))
        return value

    quarter_no = text("Quarter No", 20).upper()
    sector = text("Sector", 100)
    location = text("Location / Block", 150)
    floor = text("Floor", 20)

    type_text = _clean_text(values["Type"])
    quarter_type = TYPE_LOOKUP.get(type_text.lower())
    if quarter_type is None:
        allowed = ", ".join(label for _, label in Quarter.QuarterType.choices)
        errors.append(RowError(row_no, "Type", "is required" if not type_text
                               else f"must be one of {allowed} (got '{type_text}')"))

    area = check("Area (sq ft)", _parse_positive_decimal(values["Area (sq ft)"], max_digits=8))
    base_rate = check("Base Rate (Rs.)", _parse_positive_decimal(values["Base Rate (Rs.)"], max_digits=12))
    emd = check("EMD Amount (Rs.)", _parse_positive_decimal(values["EMD Amount (Rs.)"], max_digits=12))
    bidding_date = check("Bidding Date", _parse_date(values["Bidding Date"]))
    start_clock = check("Start Time", _parse_time(values["Start Time"]))
    duration = check("Duration (min)", _parse_duration(values["Duration (min)"]))

    start_time = None
    if bidding_date and start_clock:
        start_time = timezone.make_aware(datetime.combine(bidding_date, start_clock),
                                         timezone.get_default_timezone())
        if start_time <= now:
            errors.append(RowError(row_no, "Start Time",
                                   f"must be in the future (got {start_time:%d-%m-%Y %H:%M} IST)"))

    if errors:
        return quarter_no, None, errors
    quarter = Quarter(quarter_no=quarter_no, sector=sector, location_block=location,
                      quarter_type=quarter_type, floor=floor, area_sqft=area,
                      base_rate=base_rate, emd_amount=emd, status=Quarter.Status.VACANT)
    return quarter_no, ValidRow(quarter, start_time, duration), []
