"""
Generate the sample Excel files used for import tests and demos.

    python sample_data/make_samples.py

ALL DATA IN THESE FILES IS FAKE TEST DATA - placeholder figures, not official
SAIL data. The importer reads only the first sheet; a second "README" sheet in
each file says the data is fake.

The importer refuses auctions that start in the past, so the bidding dates are
written RELATIVE TO TODAY (a week from now). Re-run this script before a demo
if the files are older than that. The tests build the same workbooks in memory
(build_* functions below), so they never go stale.
"""
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).resolve().parent

HEADERS = [
    "Quarter No", "Sector", "Location / Block", "Type", "Floor", "Area (sq ft)",
    "Base Rate (Rs.)", "EMD Amount (Rs.)", "Bidding Date", "Start Time", "Duration (min)",
]

DAYS_AHEAD = 7  # bidding day = one week after the file is generated


def valid_rows(day):
    """8 valid rows; auctions back to back from 10:00, two of them at the same time."""
    return [
        ["B-214", "Sector 4", "Near market", "Type B", "2nd", 750, 45000, 9000, day, time(10, 0), 5],
        ["B-118", "Sector 4", "Near market", "Type B", "Ground", 750, 48000, 9500, day, time(10, 0), 5],
        ["C-305", "Sector 2", "Opposite park", "Type C", "3rd", 620, 36000, 7000, day, time(10, 10), 5],
        ["C-112", "Sector 2", "Opposite park", "Type C", "1st", 620, 39000, 7500, day, time(10, 20), 5],
        ["D-407", "Sector 9", "Near school", "Type D", "4th", 480, 24000, 5000, day, time(10, 30), None],
        ["D-203", "Sector 9", "Near school", "Type D", "2nd", 480, 27000, 5500, day, time(10, 40), 5],
        ["A-021", "Sector 12", "Main road side", "Type A", "Ground", 980, 62000, 12000, day, time(11, 0), 10],
        ["A-134", "Sector 12", "Main road side", "Type A", "1st", 980, 60000, 12000, day, time(11, 15), 10],
    ]


def invalid_rows(day):
    """Deliberate errors. Row numbers below are Excel row numbers (row 1 = header)."""
    past = date.today() - timedelta(days=1)
    return [
        ["B-214", "Sector 4", "Near market", "Type B", "2nd", 750, 45000, 9000, day, time(10, 0), 5],      # 2: valid
        ["B-214", "Sector 4", "Near market", "Type B", "3rd", 750, 46000, 9000, day, time(10, 5), 5],      # 3: duplicate Quarter No
        ["C-305", "Sector 2", "Opposite park", "Type C", "3rd", "abc", 36000, 7000, day, time(10, 10), 5],  # 4: non-numeric Area
        ["D-407", "Sector 9", "Near school", "Type D", "4th", 480, -24000, 5000, day, time(10, 20), 5],    # 5: negative Base Rate
        [None, "Sector 12", "Main road side", "Type A", "Ground", 980, 62000, 12000, day, time(10, 30), 5],  # 6: missing Quarter No
        ["A-134", "Sector 12", "Main road side", "Type A", "1st", 980, 60000, 0, day, time(10, 40), 5],     # 7: EMD not positive
        ["C-112", "Sector 2", "Opposite park", "Type C", "1st", 620, 39000, 7500, "31-02-2026", "10:50", 5],  # 8: impossible date
        ["D-203", "Sector 9", "Near school", "Type D", "2nd", 480, 27000, 5500, past, time(11, 0), 5],     # 9: start in the past
        ["B-118", "Sector 4", "Near market", "Type B", "Ground", 750, 48000, 9500, day, "25:00", 0],       # 10: bad time + zero duration
    ]


def edge_case_rows(day):
    """
    All VALID, but written the awkward ways real spreadsheets are: text dates
    and times, thousands separators, extra spaces, lower-case type, blank
    duration (uses the default), a blank row in the middle.
    """
    text_day = day.strftime("%d-%m-%Y")
    return [
        ["  e-101 ", " Sector 5 ", "Behind  hospital", "type a", "Ground", "1,050", "75,000", "15,000",
         text_day, "09:30", "10"],
        ["E-102", "Sector 5", "Behind hospital", "A", "1st", 1050.5, "72,500.50", 14500, day.strftime("%d/%m/%Y"),
         "2:15 PM", None],
        [None, None, None, None, None, None, None, None, None, None, None],          # blank row: skipped
        ["E-103", "Sector 5", "Behind hospital", "Type A", "2nd", 1050, 70000, 14000,
         datetime.combine(day, time(0, 0)), time(15, 0), 7],
        ["E-104", "Sector 5", "Behind hospital", "Type A", "3rd", 1050, 68000, 13500, day.isoformat(), "16:45:00", ""],
    ]


def build_workbook(headers, rows, note):
    wb = Workbook()
    ws = wb.active
    ws.title = "Quarters (FAKE DATA)"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        if isinstance(row[8].value, (date, datetime)):
            row[8].number_format = "DD-MM-YYYY"   # real Excel date cell
        if isinstance(row[9].value, time):
            row[9].number_format = "HH:MM"        # real Excel time cell
    for col in "ABCDEFGHIJK":
        ws.column_dimensions[col].width = 17

    readme = wb.create_sheet("README")
    readme.append(["FAKE TEST DATA - placeholder figures, NOT official SAIL data."])
    readme.append([note])
    wb.properties.title = "FAKE test data - SAIL quarter auction training project"
    return wb


def build_valid(day):
    return build_workbook(HEADERS, valid_rows(day), "8 valid rows; should import cleanly.")


def build_invalid(day):
    return build_workbook(HEADERS, invalid_rows(day),
                          "Deliberate errors: duplicate Quarter No (row 3), non-numeric Area (row 4), "
                          "negative Base Rate (row 5), missing Quarter No (row 6), EMD 0 (row 7), "
                          "impossible date (row 8), start in the past (row 9), "
                          "bad time and zero duration (row 10).")


def build_bad_header(day):
    bad = ["Quarter Number"] + HEADERS[1:]  # "Quarter Number" instead of "Quarter No"
    return build_workbook(bad, valid_rows(day), 'Header error: column A is "Quarter Number".')


def build_edge_cases(day):
    return build_workbook(HEADERS, edge_case_rows(day),
                          "Valid but awkward: text dates/times, thousands separators, extra spaces, "
                          "lower-case type, blank duration, a blank row. Should import 4 quarters.")


def bidding_day():
    return date.today() + timedelta(days=DAYS_AHEAD)


if __name__ == "__main__":
    day = bidding_day()
    for name, builder in [("quarters_sample.xlsx", build_valid), ("quarters_invalid.xlsx", build_invalid),
                          ("quarters_bad_header.xlsx", build_bad_header),
                          ("quarters_edge_cases.xlsx", build_edge_cases)]:
        builder(day).save(HERE / name)
        print(f"wrote {HERE / name}")
    print(f"Bidding date in the files: {day:%d-%m-%Y}")
